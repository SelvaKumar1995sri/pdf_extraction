"""
converter.py — Convert extracted HTML to DOCX, HTML, or EPUB.
"""

import base64
import io
import os
import re
import tempfile

from bs4 import BeautifulSoup, NavigableString, Tag


# ── Shared book CSS ───────────────────────────────────────────────────────────

BOOK_CSS = """
body { font-family: Georgia, 'Times New Roman', serif; font-size: 11pt;
       line-height: 1.65; color: #1a1a1a; margin: 2em 3em; }
h1 { font-size: 1.9em; font-weight: bold; text-align: center;
     margin: 1.2em 0 0.6em; }
h2 { font-size: 1.45em; font-weight: bold; text-align: center;
     margin: 1em 0 0.5em; }
h3 { font-size: 1.18em; font-weight: bold; margin: 0.8em 0 0.4em; }
p  { margin: 0 0 0.55em 0; }
figure { text-align: center; margin: 1.5em 0; }
figure img { max-width: 100%; height: auto; }
"""


# ── Helpers ───────────────────────────────────────────────────────────────────

def _parse_pages(html: str) -> list:
    soup  = BeautifulSoup(html, "html.parser")
    pages = soup.find_all("div", class_="pdf-page")
    return pages if pages else [soup]


def _img_src_to_bytes(src: str) -> bytes | None:
    m = re.match(r"data:image/[^;]+;base64,(.+)", src, re.DOTALL)
    return base64.b64decode(m.group(1)) if m else None


def _img_ext(src: str) -> str:
    m = re.match(r"data:image/([^;]+);", src)
    return m.group(1) if m else "png"


# ── HTML export ───────────────────────────────────────────────────────────────

def to_html(html: str, title: str = "Extracted Book") -> bytes:
    soup  = BeautifulSoup(html, "html.parser")
    pages = _parse_pages(html)

    font_styles = "\n".join(str(s) for s in soup.find_all("style"))
    bodies = []
    for page in pages:
        inner = "".join(str(c) for c in page.children)
        bodies.append(f'<div class="pdf-page">{inner}</div>')

    doc = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8"/>
<title>{title}</title>
{font_styles}
<style>
{BOOK_CSS}
.pdf-page {{ max-width:780px; margin:0 auto 2em; padding:3em 4em;
             background:#fff; box-shadow:0 2px 12px rgba(0,0,0,.12); }}
</style>
</head>
<body>
{"".join(bodies)}
</body>
</html>"""
    return doc.encode("utf-8")


# ── EPUB export ───────────────────────────────────────────────────────────────

def _build_epub_zip(chapters: list[dict], images: dict[str, bytes],
                    title: str, author: str) -> bytes:
    """Build a valid EPUB 3 ZIP archive directly — no ebooklib needed.

    chapters  : list of {"title": str, "filename": str, "xhtml": str}
    images    : {"images/img_0001.png": <bytes>, ...}
    """
    import zipfile, uuid

    uid = str(uuid.uuid4())

    # ── content.opf ────────────────────────────────────────────────────
    manifest_items = [
        '<item id="style" href="style.css" media-type="text/css"/>',
        '<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" '
        'properties="nav"/>',
    ]
    spine_items = []
    for ch in chapters:
        cid = ch["filename"].replace(".xhtml", "").replace("-", "_")
        manifest_items.append(
            f'<item id="{cid}" href="{ch["filename"]}" '
            f'media-type="application/xhtml+xml"/>'
        )
        spine_items.append(f'<itemref idref="{cid}"/>')
    for img_path in images:
        fname = img_path.split("/")[-1]
        ext   = fname.rsplit(".", 1)[-1].lower()
        mime  = {"jpg": "image/jpeg", "jpeg": "image/jpeg",
                 "png": "image/png",  "gif":  "image/gif",
                 "svg": "image/svg+xml"}.get(ext, "image/png")
        iid   = img_path.replace("/", "_").replace(".", "_")
        manifest_items.append(
            f'<item id="{iid}" href="{img_path}" media-type="{mime}"/>'
        )

    opf = f"""<?xml version="1.0" encoding="utf-8"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0"
         unique-identifier="bookid" xml:lang="en">
  <metadata xmlns:dc="http://purl.org/dc/elements/1.1/">
    <dc:identifier id="bookid">{uid}</dc:identifier>
    <dc:title>{title}</dc:title>
    <dc:creator>{author}</dc:creator>
    <dc:language>en</dc:language>
  </metadata>
  <manifest>
    {"".join(manifest_items)}
  </manifest>
  <spine>
    {"".join(spine_items)}
  </spine>
</package>"""

    # ── nav.xhtml ───────────────────────────────────────────────────────
    nav_items = "\n".join(
        f'<li><a href="{ch["filename"]}">{ch["title"]}</a></li>'
        for ch in chapters
    )
    nav = f"""<?xml version="1.0" encoding="utf-8"?>
<!DOCTYPE html>
<html xmlns="http://www.w3.org/1999/xhtml"
      xmlns:epub="http://www.idpf.org/2007/ops" xml:lang="en">
<head><meta charset="utf-8"/><title>Table of Contents</title></head>
<body>
  <nav epub:type="toc" id="toc">
    <h1>Contents</h1>
    <ol>{nav_items}</ol>
  </nav>
</body>
</html>"""

    # ── Assemble ZIP ────────────────────────────────────────────────────
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        # mimetype MUST be first and uncompressed
        zi = zipfile.ZipInfo("mimetype")
        zi.compress_type = zipfile.ZIP_STORED
        zf.writestr(zi, "application/epub+zip")

        zf.writestr("META-INF/container.xml",
            '<?xml version="1.0"?>'
            '<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container"'
            ' version="1.0">'
            '<rootfiles>'
            '<rootfile full-path="OEBPS/content.opf"'
            ' media-type="application/oebps-package+xml"/>'
            '</rootfiles></container>')

        zf.writestr("OEBPS/content.opf",  opf)
        zf.writestr("OEBPS/nav.xhtml",    nav)
        zf.writestr("OEBPS/style.css",    BOOK_CSS)

        for ch in chapters:
            zf.writestr(f"OEBPS/{ch['filename']}", ch["xhtml"])

        for img_path, img_bytes in images.items():
            zf.writestr(f"OEBPS/{img_path}", img_bytes)

    buf.seek(0)
    return buf.getvalue()


def to_epub(html: str, title: str = "Extracted Book",
            author: str = "Unknown") -> bytes:
    """Build EPUB 3 directly as a ZIP — no ebooklib required."""

    soup  = BeautifulSoup(html, "html.parser")
    pages = soup.find_all("div", class_="pdf-page") or [soup]

    chapters: list[dict] = []
    images:   dict[str, bytes] = {}
    img_idx = 0

    for pi, page in enumerate(pages, 1):
        # Extract and replace base64 images with relative paths
        for img_tag in page.find_all("img"):
            src = img_tag.get("src", "")
            if not src.startswith("data:image"):
                continue
            img_bytes = _img_src_to_bytes(src)
            if not img_bytes:
                continue
            ext      = _img_ext(src)
            img_path = f"images/img_{img_idx:04d}.{ext}"
            images[img_path] = img_bytes
            img_tag["src"]   = img_path
            img_idx += 1

        # Chapter title
        heading    = page.find(re.compile(r"^h[123]$"))
        chap_title = heading.get_text(strip=True) if heading else f"Page {pi}"

        # Body — strip <style> blocks
        body = "".join(
            str(c) for c in page.children
            if not (isinstance(c, Tag) and c.name == "style")
        ).strip() or f"<p>{chap_title}</p>"

        xhtml = (
            "<?xml version='1.0' encoding='utf-8'?>"
            "<!DOCTYPE html>"
            "<html xmlns='http://www.w3.org/1999/xhtml' xml:lang='en'>"
            "<head>"
            f"<meta charset='utf-8'/>"
            f"<title>{chap_title}</title>"
            "<link rel='stylesheet' type='text/css' href='style.css'/>"
            "</head>"
            f"<body>{body}</body>"
            "</html>"
        )
        chapters.append({
            "title":    chap_title,
            "filename": f"chap_{pi:04d}.xhtml",
            "xhtml":    xhtml,
        })

    if not chapters:
        raise ValueError("No content found to build EPUB.")

    return _build_epub_zip(chapters, images, title, author)


# ── DOCX export ───────────────────────────────────────────────────────────────

def to_docx(html: str, title: str = "Extracted Book") -> bytes:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Default Normal style
    normal = doc.styles["Normal"]
    normal.font.name = "Georgia"
    normal.font.size = Pt(11)

    def _process_children(para, el):
        """Walk element children and add styled runs."""
        for child in el.children:
            if isinstance(child, NavigableString):
                text = str(child)
                if text:
                    para.add_run(text)
            elif isinstance(child, Tag):
                if child.name == "strong":
                    run = para.add_run(child.get_text())
                    run.bold = True
                elif child.name == "em":
                    run = para.add_run(child.get_text())
                    run.italic = True
                elif child.name == "span":
                    style_attr = child.get("style", "")
                    # Recurse to pick up <strong>/<em> inside span
                    _process_children(para, child)
                else:
                    _process_children(para, child)

    pages = _parse_pages(html)

    for pi, page in enumerate(pages):
        if pi > 0:
            doc.add_page_break()

        for element in page.children:
            if not isinstance(element, Tag):
                continue

            tag = element.name

            # Headings
            if tag in ("h1", "h2", "h3"):
                level   = int(tag[1])
                heading = doc.add_heading("", level=level)
                heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _process_children(heading, element)
                continue

            # Paragraphs
            if tag == "p":
                style_attr = element.get("style", "")

                # TOC entry (flex, space-between)
                if "space-between" in style_attr:
                    spans = [s for s in element.find_all("span", recursive=False)
                             if not s.get("style", "").startswith("flex")]
                    left  = spans[0].get_text()  if len(spans) > 0 else element.get_text()
                    right = spans[-1].get_text() if len(spans) > 1 else ""
                    para  = doc.add_paragraph(style="Normal")
                    para.add_run(left)
                    para.add_run("\t" + right)
                    continue

                # Alignment
                align_str = re.search(r"text-align:(\w+)", style_attr)
                align_map = {
                    "center": WD_ALIGN_PARAGRAPH.CENTER,
                    "right":  WD_ALIGN_PARAGRAPH.RIGHT,
                }
                para = doc.add_paragraph(style="Normal")
                if align_str:
                    para.alignment = align_map.get(
                        align_str.group(1), WD_ALIGN_PARAGRAPH.LEFT
                    )
                _process_children(para, element)
                continue

            # Images
            if tag == "figure":
                img = element.find("img")
                if img:
                    src       = img.get("src", "")
                    img_bytes = _img_src_to_bytes(src)
                    if img_bytes:
                        try:
                            para      = doc.add_paragraph()
                            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            run       = para.add_run()
                            run.add_picture(io.BytesIO(img_bytes), width=Inches(4))
                        except Exception:
                            pass
                continue

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
